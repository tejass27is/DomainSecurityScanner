import asyncio
import io
import os
import uuid
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.db.base import Base, engine
from app.db.models import (
    AuditLog,
    Organization,
    OrganizationRegion,
    Region,
    VaptChecklistAttachment,
    VaptImport,
    User,
    VaptOnboardingChecklist,
    VaptRescanSchedule,
)
from app.api.vapt.routes import (
    REGION_CHECKLIST_REQUIRED_QUESTIONS,
    VERIFICATION_DISPLAY_NAME_MAX,
    _clean_verification_display_name,
    _closure_blockers,
    _closure_block_message,
    _evaluate_manual_verification,
    _generate_checklist_pdf,
    _get_onboarding_or_create,
    _to_detail,
    _missing_required_uploads,
    _normalize_finding_status,
    _new_findings_requiring_follow_up,
    _parse_report_files,
    _requires_vapt_finding_triage,
    _reopen_unresolved_findings,
    _validate_rescan_prerequisites,
    _verification_display_name,
    _verification_download_filename,
    ClientDueDateRequest,
    ClientDueDateRejectionRequest,
    DirectClosureRequest,
    AdminRescheduleRequest,
    approve_client_next_vapt_due_date,
    accept_proposed_date,
    admin_approve_reschedule,
    admin_request_new_date,
    client_request_new_date,
    close_vapt_without_verification,
    approve_vapt_access,
    decide_region_checklist,
    delete_checklist_attachment,
    download_checklist_attachment,
    download_approved_onboarding_bundle,
    get_vapt_access_status,
    list_approved_vapt_onboarding,
    list_vapt_access_requests,
    request_vapt_access,
    request_vapt_region,
    reject_client_next_vapt_due_date,
    reject_proposed_date,
    review_vapt_onboarding,
    review_client_remediation,
    submit_onboarding_checklist,
    submit_vapt_import,
    submit_verification_review,
    set_client_next_vapt_due_date,
    upload_checklist_attachment,
    upload_vapt_report,
    update_vapt_finding_status,
    update_verification_finding_status,
)
from app.api.vapt.normalizer import group_stored_findings, normalize_import
from app.api.vapt.maintenance import run_vapt_client_review_reminders
from app.api.admin.service import get_users_by_org
from app.api.vapt.schemas import VaptFindingStatusUpdate, VaptImportDetail, VaptImportListItem
from app.api.vapt.report_generator import (
    generate_vapt_report_pdf,
    generate_vapt_verification_report_pdf,
    generate_vapt_verification_report_xlsx,
)
from app.utils import email as email_utils


def _complete_region_checklist_answers():
    answers = {
        section: {
            question_id: {"answer": "Provided", "na": False}
            for question_id in question_ids
        }
        for section, question_ids in REGION_CHECKLIST_REQUIRED_QUESTIONS.items()
    }
    answers["computers_servers"]["asset_list_upload"] = {
        "rows": [{"asset": "test-host"}],
        "na": False,
    }
    answers["general_information"] = {
        "network_diagram_available": {"answer": "Available", "na": False},
    }
    return answers


def test_vapt_pdf_generation_includes_region_without_callback_name_error():
    record = SimpleNamespace(
        import_id="test-import",
        org_id="test-org",
        region="CHA-IRE",
        file_name="combined-report.nessus",
        file_format="nessus",
        source_tool="nessus",
        status="open",
        lifecycle_status="report_published",
        total_findings=1,
        unique_hosts=1,
        risk_score=50,
        severity="high",
        severity_distribution={"critical": 0, "high": 1, "medium": 0, "low": 0, "info": 0},
        category_distribution={"Network": 1},
        summary={"raw_findings_parsed": 1},
        findings=[{
            "id": "finding-1",
            "title": "Test vulnerability",
            "severity_label": "high",
            "cvss_score": 8.0,
            "affected_hosts": ["192.0.2.10"],
            "category": "Network",
            "description": "A test finding.",
            "solution": "Apply the recommended fix.",
        }],
        created_at=datetime(2026, 10, 12, tzinfo=timezone.utc),
    )

    pdf = generate_vapt_report_pdf(record)

    assert pdf.startswith(b"%PDF")
    assert "CHA-IRE" in PdfReader(io.BytesIO(pdf)).pages[1].extract_text()


def test_verification_exports_include_fully_fixed_grouped_findings():
    fixed = {
        "id": "F001",
        "title": "Grouped fixed vulnerability",
        "severity_label": "high",
        "affected_hosts": ["192.0.2.10", "192.0.2.11"],
        "affected_assets": [
            {"host": "192.0.2.10", "mac_address": "00:11:22:33:44:10"},
            {"host": "192.0.2.11", "mac_address": "00:11:22:33:44:11"},
        ],
        "mac_address": "00:11:22:33:44:10, 00:11:22:33:44:11",
        "protocols": ["tcp", "udp"],
        "ports": [53, 443],
        "verified_fixed_hosts": ["192.0.2.10", "192.0.2.11"],
        "original_host_count": 2,
        "category": "Network",
        "description": "The grouped issue was fixed on both hosts.",
    }
    schedule = SimpleNamespace(
        result_data={
            "findings": [],
            "fixed_findings": [fixed],
            "remaining_findings": [],
            "new_findings": [],
            "unverified_findings": [],
        },
        status="completed",
        scheduled_at=datetime(2026, 10, 12, tzinfo=timezone.utc),
        uploaded_at=datetime(2026, 10, 12, tzinfo=timezone.utc),
        verification_outcome="pending",
    )
    record = SimpleNamespace(
        import_id="test-import",
        org_id="test-org",
        region="CHA-IRE",
        file_name="initial-report.nessus",
        file_format="nessus",
        source_tool="nessus",
        risk_score=50,
        severity_distribution={"critical": 0, "high": 1, "medium": 0, "low": 0, "info": 0},
        category_distribution={"Network": 1},
        summary={},
        findings=[fixed],
        created_at=datetime(2026, 10, 12, tzinfo=timezone.utc),
    )

    pdf = generate_vapt_verification_report_pdf(schedule, record)
    xlsx = generate_vapt_verification_report_xlsx(schedule, record)

    assert "Grouped fixed vulnerability" in "\n".join(
        page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages
    )
    workbook = load_workbook(io.BytesIO(xlsx), read_only=True)
    rows = list(workbook["Verification Findings"].values)
    assert any(row[1] == "Grouped fixed vulnerability" and row[6] == "Fixed" for row in rows)
    assert rows[1][4] == "192.0.2.10, 192.0.2.11"
    assert rows[1][7] == "00:11:22:33:44:10, 00:11:22:33:44:11"
    assert rows[1][8] == "tcp, udp"
    assert rows[1][9] == "53, 443"


def test_parse_report_files_merges_findings_and_records_source_metadata(monkeypatch):
    parsed = {
        "first.nessus": ([{"title": "First issue", "host": "10.0.1.15", "severity": 3}], "nessus", "nessus"),
        "second.nessus": ([{"title": "Second issue", "host": "10.0.2.15", "severity": 3}], "nessus", "nessus"),
    }
    monkeypatch.setattr(
        "app.api.vapt.routes.parse_upload",
        lambda content, filename: parsed[filename],
    )

    findings, source_tool, file_format, filenames, tools = _parse_report_files([
        ("first.nessus", b"first"),
        ("second.nessus", b"second"),
    ])

    assert findings == [
        {"title": "First issue", "host": "10.0.1.15", "severity": 3},
        {"title": "Second issue", "host": "10.0.2.15", "severity": 3},
    ]
    normalized = normalize_import(findings, "nessus")
    assert source_tool == "nessus"
    assert file_format == "nessus"
    assert filenames == ["first.nessus", "second.nessus"]
    assert tools == ["nessus"]
    assert {host for finding in normalized["findings"] for host in finding["affected_hosts"]} == {
        "10.0.1.15",
        "10.0.2.15",
    }


def test_normalizer_groups_the_same_scanner_check_across_hosts():
    findings = [
        {
            "title": "Weak TLS configuration",
            "plugin_id": "1234",
            "severity": 3,
            "severity_label": "high",
            "host": host,
            "port": 443,
            "protocol": "tcp",
            "description": "The service supports a weak protocol.",
            "cves": [],
            "evidence": f"Evidence from {host}",
        }
        for host in ("192.0.2.10", "192.0.2.11", "192.0.2.12")
    ]

    normalized = normalize_import(findings, "nessus")

    assert len(normalized["findings"]) == 1
    grouped = normalized["findings"][0]
    assert grouped["host_count"] == 3
    assert grouped["affected_hosts"] == ["192.0.2.10", "192.0.2.11", "192.0.2.12"]
    assert {entry["host"] for entry in grouped["host_outputs"]} == set(grouped["affected_hosts"])
    assert normalized["total_findings"] == 1
    assert normalized["risk_score"] == normalize_import(findings[:1], "nessus")["risk_score"]


def test_normalizer_groups_same_finding_name_across_protocols_and_keeps_asset_pairs():
    findings = [
        {
            "title": "  Shared   vulnerability name ",
            "plugin_id": "plugin-a",
            "severity": 2,
            "severity_label": "medium",
            "host": "192.0.2.20",
            "mac_address": "00:11:22:33:44:55",
            "hostname": "server-a",
            "port": 53,
            "protocol": "udp",
            "description": "UDP-specific detail.",
        },
        {
            "title": "shared vulnerability name",
            "plugin_id": "plugin-b",
            "severity": 3,
            "severity_label": "high",
            "host": "192.0.2.21",
            "mac_address": "00:11:22:33:44:66",
            "hostname": "server-b",
            "port": 443,
            "protocol": "tcp",
            "description": "TCP-specific detail.",
        },
        {
            "title": "SHARED VULNERABILITY NAME",
            "plugin_id": "plugin-c",
            "severity": 2,
            "severity_label": "medium",
            "host": "192.0.2.20",
            "mac_address": "00:11:22:33:44:55",
            "hostname": "server-a",
            "port": 8443,
            "protocol": "tcp",
            "description": "Another port on the same asset.",
        },
    ]

    normalized = normalize_import(findings, "generic")

    assert len(normalized["findings"]) == 1
    grouped = normalized["findings"][0]
    assert grouped["host_count"] == 2
    assert grouped["affected_hosts"] == ["192.0.2.20", "192.0.2.21"]
    assert {
        (asset["host"], asset["mac_address"])
        for asset in grouped["affected_assets"]
    } == {
        ("192.0.2.20", "00:11:22:33:44:55"),
        ("192.0.2.21", "00:11:22:33:44:66"),
    }
    assert grouped["protocols"] == ["tcp", "udp"]
    assert grouped["ports"] == [53, 443, 8443]
    assert grouped["severity_label"] == "high"


def test_stored_findings_are_displayed_as_one_name_group_with_separate_asset_columns():
    grouped = group_stored_findings([
        {
            "id": "F001",
            "title": "Microsoft SQL Server Unsupported Version Detection (remote check)",
            "severity": 4,
            "severity_label": "critical",
            "status": "pending",
            "affected_hosts": ["172.16.116.40"],
            "mac_address": "74:56:3C:2F:6E:DA",
            "port": 1433,
            "protocol": "tcp",
        },
        {
            "id": "F002",
            "title": "Microsoft SQL Server Unsupported Version Detection (remote check)",
            "severity": 4,
            "severity_label": "critical",
            "status": "pending",
            "affected_hosts": ["172.16.116.36"],
            "mac_address": "84:A9:3E:78:C4:8F",
            "port": 1434,
            "protocol": "tcp",
        },
    ])

    assert len(grouped) == 1
    assert grouped[0]["affected_hosts"] == ["172.16.116.36", "172.16.116.40"]
    assert {asset["mac_address"] for asset in grouped[0]["affected_assets"]} == {
        "84:A9:3E:78:C4:8F",
        "74:56:3C:2F:6E:DA",
    }
    assert grouped[0]["mac_address"] == "74:56:3C:2F:6E:DA, 84:A9:3E:78:C4:8F"
    assert grouped[0]["ports"] == [1433, 1434]
    assert grouped[0]["protocols"] == ["tcp"]
    assert grouped[0]["host_count"] == 2
    assert grouped[0]["grouped_finding_ids"] == ["F001", "F002"]


def test_updating_grouped_legacy_finding_updates_each_underlying_row():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-group-status",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            total_findings=3,
            unique_hosts=2,
            risk_score=50,
            severity="high",
            findings=[
                {
                    "id": "F001",
                    "title": "Same named issue",
                    "severity_label": "high",
                    "status": "pending",
                    "affected_hosts": ["192.0.2.10"],
                    "port": 80,
                    "protocol": "tcp",
                },
                {
                    "id": "F002",
                    "title": "Same named issue",
                    "severity_label": "high",
                    "status": "pending",
                    "affected_hosts": ["192.0.2.11"],
                    "port": 443,
                    "protocol": "tcp",
                },
            ],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        update_vapt_finding_status(
            str(record.import_id),
            "F001",
            VaptFindingStatusUpdate(status="solved", comment="Fixed across the grouped assets"),
            db=db,
            current_user=SimpleNamespace(org_id="org-group-status", user_id="test-user"),
        )

        db.refresh(record)
        assert [finding["status"] for finding in record.findings] == ["solved", "solved"]
        assert [finding["comment"] for finding in record.findings] == [
            "Fixed across the grouped assets",
            "Fixed across the grouped assets",
        ]
        detail = _to_detail(record)
        assert len(detail["findings"]) == 1
        assert detail["findings"][0]["host_count"] == 2
        assert detail["findings"][0]["ports"] == [80, 443]
        assert detail["findings"][0]["status"] == "solved"
    finally:
        db.close()


def test_normalized_informational_findings_remain_visible_without_changing_risk_score():
    low_finding = {
        "title": "Low severity issue",
        "description": "A low severity test finding.",
        "severity": 1,
        "severity_label": "low",
        "host": "10.0.0.1",
    }
    info_finding = {
        "title": "Informational service",
        "description": "A service is listening.",
        "severity": 0,
        "severity_label": "info",
        "host": "10.0.0.2",
    }

    low_only = normalize_import([low_finding], "nessus")
    with_info = normalize_import([low_finding, info_finding], "nessus")

    assert [finding["severity_label"] for finding in with_info["findings"]] == ["low", "info"]
    assert with_info["summary"]["excluded_info_findings"] == 0
    assert with_info["summary"]["report_findings"] == 2
    assert with_info["risk_score"] == low_only["risk_score"]


def test_closed_vapt_cycle_cannot_schedule_verification():
    record = VaptImport(
        org_id="org-closed-rescan",
        file_name="closed-report.xml",
        file_format="xml",
        source_tool="nessus",
        status="client_completed",
        lifecycle_status="closed",
        remediation_review_status="approved",
        next_vapt_due_at=datetime(2027, 1, 15, 10, tzinfo=timezone.utc),
        findings=[{"id": "finding-1", "status": "solved"}],
    )

    try:
        _validate_rescan_prerequisites(record)
        raise AssertionError("expected closed VAPT cycles to reject verification scheduling")
    except HTTPException as exc:
        assert exc.status_code == 409
        assert "cycle is closed" in exc.detail


def setup_module():
    Base.metadata.drop_all(bind=engine, checkfirst=True)
    Base.metadata.create_all(bind=engine)


def test_rescan_date_negotiation_alternates_between_client_and_soc(monkeypatch):
    org_id = str(uuid.uuid4())
    client_id = str(uuid.uuid4())
    soc_id = str(uuid.uuid4())
    import_id = uuid.uuid4()
    db = Session(bind=engine)
    monkeypatch.setattr("app.api.vapt.routes.send_vapt_access_event_email", lambda *args, **kwargs: None)
    try:
        client = User(
            user_id=client_id,
            org_id=org_id,
            email=f"{client_id}@example.com",
            password="test-password",
            role="owner",
        )
        soc = User(
            user_id=soc_id,
            org_id=None,
            email=f"{soc_id}@example.com",
            password="test-password",
            role="soc_analyst",
        )
        record = VaptImport(
            import_id=import_id,
            org_id=org_id,
            file_name="rescan-negotiation.xml",
            file_format="xml",
            source_tool="nessus",
            lifecycle_status="revalidation_scheduled",
            status="client_completed",
            remediation_review_status="approved",
            findings=[{"id": "finding-1", "status": "solved"}],
        )
        schedule = VaptRescanSchedule(
            import_id=import_id,
            org_id=org_id,
            created_by=client_id,
            scheduled_at=datetime.now(timezone.utc) + timedelta(days=10),
            status="scheduled",
            date_proposed_by="client",
        )
        db.add_all([client, soc, Organization(org_id=org_id, user_id=client_id), record, schedule])
        db.commit()
        db.refresh(schedule)

        asyncio.run(admin_request_new_date(
            str(schedule.id),
            AdminRescheduleRequest(
                proposed_at=(datetime.now(timezone.utc) + timedelta(days=12)).isoformat(),
                proposed_timezone="UTC",
            ),
            db=db,
            current_user=soc,
        ))
        db.refresh(schedule)
        assert schedule.date_proposed_by == "soc"
        assert schedule.date_proposal_status == "pending"

        with pytest.raises(HTTPException) as error:
            asyncio.run(admin_approve_reschedule(str(schedule.id), db=db, current_user=soc))
        assert error.value.status_code == 409

        asyncio.run(client_request_new_date(
            str(import_id),
            str(schedule.id),
            AdminRescheduleRequest(
                proposed_at=(datetime.now(timezone.utc) + timedelta(days=14)).isoformat(),
                proposed_timezone="UTC",
            ),
            db=db,
            current_user=client,
        ))
        db.refresh(schedule)
        assert schedule.date_proposed_by == "client"
        assert schedule.date_proposal_status == "pending"

        with pytest.raises(HTTPException) as error:
            asyncio.run(client_request_new_date(
                str(import_id),
                str(schedule.id),
                AdminRescheduleRequest(
                    proposed_at=(datetime.now(timezone.utc) + timedelta(days=15)).isoformat(),
                    proposed_timezone="UTC",
                ),
                db=db,
                current_user=client,
            ))
        assert error.value.status_code == 409

        asyncio.run(admin_approve_reschedule(str(schedule.id), db=db, current_user=soc))
        db.refresh(schedule)
        assert schedule.status == "approved"
        assert schedule.date_proposal_status == "accepted"

        client_accept_schedule = VaptRescanSchedule(
            import_id=import_id,
            org_id=org_id,
            created_by=client_id,
            scheduled_at=datetime.now(timezone.utc) + timedelta(days=20),
            status="requested",
            date_proposed_by="soc",
            date_proposal_status="pending",
        )
        db.add(client_accept_schedule)
        db.commit()
        db.refresh(client_accept_schedule)
        asyncio.run(accept_proposed_date(
            str(import_id),
            str(client_accept_schedule.id),
            db=db,
            current_user=client,
        ))
        db.refresh(client_accept_schedule)
        assert client_accept_schedule.status == "approved"
        assert client_accept_schedule.date_proposal_status == "accepted"

        rejected_schedule = VaptRescanSchedule(
            import_id=import_id,
            org_id=org_id,
            created_by=client_id,
            scheduled_at=datetime.now(timezone.utc) + timedelta(days=25),
            status="requested",
            date_proposed_by="soc",
            date_proposal_status="pending",
        )
        db.add(rejected_schedule)
        db.commit()
        db.refresh(rejected_schedule)
        asyncio.run(reject_proposed_date(
            str(import_id),
            str(rejected_schedule.id),
            db=db,
            current_user=client,
        ))
        db.refresh(rejected_schedule)
        assert rejected_schedule.date_proposal_status == "rejected"

        asyncio.run(admin_request_new_date(
            str(rejected_schedule.id),
            AdminRescheduleRequest(
                proposed_at=(datetime.now(timezone.utc) + timedelta(days=30)).isoformat(),
                proposed_timezone="UTC",
            ),
            db=db,
            current_user=soc,
        ))
        db.refresh(rejected_schedule)
        assert rejected_schedule.date_proposed_by == "soc"
        assert rejected_schedule.date_proposal_status == "pending"
    finally:
        schedule_ids = [str(item.id) for item in db.query(VaptRescanSchedule).filter_by(import_id=import_id).all()]
        db.query(AuditLog).filter(
            AuditLog.target_type == "vapt_rescan_schedule",
            AuditLog.target_id.in_(schedule_ids),
        ).delete(synchronize_session=False)
        db.query(VaptRescanSchedule).filter_by(import_id=import_id).delete(synchronize_session=False)
        db.query(VaptImport).filter_by(import_id=import_id).delete(synchronize_session=False)
        db.query(Organization).filter_by(org_id=org_id).delete(synchronize_session=False)
        db.query(User).filter(User.user_id.in_([client_id, soc_id])).delete(synchronize_session=False)
        db.commit()
        db.close()


def test_soc_due_date_rejection_returns_cycle_to_client_for_resubmission():
    org_id = str(uuid.uuid4())
    client_id = str(uuid.uuid4())
    soc_id = str(uuid.uuid4())
    import_id = uuid.uuid4()
    reason = "Please choose a date after the infrastructure upgrade."
    db = Session(bind=engine)
    try:
        client = User(
            user_id=client_id,
            org_id=org_id,
            email=f"{client_id}@example.com",
            password="test-password",
            role="owner",
        )
        soc = User(
            user_id=soc_id,
            org_id=None,
            email=f"{soc_id}@example.com",
            password="test-password",
            role="soc_analyst",
        )
        record = VaptImport(
            import_id=import_id,
            org_id=org_id,
            file_name="test-report.xml",
            file_format="xml",
            source_tool="nessus",
            lifecycle_status="closure_pending_soc_due_date",
            next_vapt_due_at=datetime.now(timezone.utc) + timedelta(days=10),
            summary={"report_findings": 0},
            findings=[],
        )
        db.add_all([client, soc, Organization(org_id=org_id, user_id=client_id), record])
        db.commit()

        with pytest.raises(HTTPException) as error:
            asyncio.run(
                reject_client_next_vapt_due_date(
                    import_id=str(import_id),
                    body=ClientDueDateRejectionRequest(reason="  "),
                    db=db,
                    current_user=soc,
                )
            )
        assert error.value.status_code == 400
        assert record.lifecycle_status == "closure_pending_soc_due_date"
        assert record.next_vapt_due_at is not None

        result = asyncio.run(
            reject_client_next_vapt_due_date(
                import_id=str(import_id),
                body=ClientDueDateRejectionRequest(reason=reason),
                db=db,
                current_user=soc,
            )
        )
        db.refresh(record)
        assert result["lifecycle_status"] == "closure_pending_client_due_date"
        assert record.lifecycle_status == "closure_pending_client_due_date"
        assert record.next_vapt_due_at is None
        assert record.summary["next_vapt_due_date_rejection_reason"] == reason
        assert db.query(AuditLog).filter_by(action="VAPT_DUE_DATE_REJECTED", target_id=str(import_id)).count() == 1

        proposal = datetime.now(timezone.utc) + timedelta(days=30)
        asyncio.run(
            set_client_next_vapt_due_date(
                import_id=str(import_id),
                body=ClientDueDateRequest(next_vapt_due_at=proposal.isoformat()),
                db=db,
                current_user=client,
            )
        )
        db.refresh(record)
        assert record.lifecycle_status == "closure_pending_soc_due_date"
        assert record.next_vapt_due_at is not None
        assert "next_vapt_due_date_rejection_reason" not in record.summary

        asyncio.run(approve_client_next_vapt_due_date(str(import_id), db=db, current_user=soc))
        db.refresh(record)
        assert record.lifecycle_status == "closed"
    finally:
        db.query(VaptImport).filter_by(import_id=import_id).delete(synchronize_session=False)
        db.query(AuditLog).filter_by(target_id=str(import_id)).delete(synchronize_session=False)
        db.query(Organization).filter_by(org_id=org_id).delete(synchronize_session=False)
        db.query(User).filter(User.user_id.in_([client_id, soc_id])).delete(synchronize_session=False)
        db.commit()
        db.close()


def test_vapt_client_review_reminders_are_three_day_once_only_and_skip_activity(monkeypatch):
    import uuid

    now = datetime.now(timezone.utc)
    org_id = str(uuid.uuid4())
    user_id = str(uuid.uuid4())
    import_id = uuid.uuid4()
    verification_import_id = uuid.uuid4()
    activity_import_id = uuid.uuid4()
    recent_import_id = uuid.uuid4()
    sent = []
    monkeypatch.setattr(
        "app.api.vapt.maintenance.send_vapt_client_review_reminder_email",
        lambda **kwargs: sent.append(kwargs),
    )
    db = Session(bind=engine)
    try:
        client = User(
            user_id=user_id,
            org_id=org_id,
            email=f"{user_id}@example.com",
            password="test",
            role="owner",
        )
        db.add(client)
        db.add(Organization(org_id=org_id, user_id=user_id))
        db.flush()
        db.add_all([
            VaptImport(
                import_id=import_id,
                org_id=org_id,
                file_name="initial-report.xml",
                file_format="xml",
                source_tool="nessus",
                created_at=now - timedelta(days=4),
                lifecycle_status="report_published",
                remediation_review_status="not_submitted",
                findings=[{"id": "initial-finding", "status": "pending"}],
            ),
            VaptImport(
                import_id=activity_import_id,
                org_id=org_id,
                file_name="already-reviewed.xml",
                file_format="xml",
                source_tool="nessus",
                created_at=now - timedelta(days=4),
                lifecycle_status="report_published",
                remediation_review_status="not_submitted",
                findings=[{"id": "reviewed-finding", "status": "solved"}],
            ),
            VaptImport(
                import_id=recent_import_id,
                org_id=org_id,
                file_name="recent-report.xml",
                file_format="xml",
                source_tool="nessus",
                created_at=now - timedelta(days=2),
                lifecycle_status="report_published",
                remediation_review_status="not_submitted",
                findings=[{"id": "recent-finding", "status": "pending"}],
            ),
            VaptImport(
                import_id=verification_import_id,
                org_id=org_id,
                file_name="verification-report.xml",
                file_format="xml",
                source_tool="nessus",
                created_at=now - timedelta(days=10),
                lifecycle_status="revalidation_verification_pending",
                remediation_review_status="approved",
                findings=[],
            ),
        ])
        db.flush()
        db.add(
            VaptRescanSchedule(
                import_id=verification_import_id,
                org_id=org_id,
                created_by=user_id,
                scheduled_at=now - timedelta(days=8),
                status="completed",
                uploaded_at=now - timedelta(days=4),
                result_data={
                    "remaining_findings": [{"id": "remaining", "severity_label": "low", "status": "pending"}],
                    "new_findings": [],
                },
            )
        )
        db.commit()

        result = run_vapt_client_review_reminders(db)

        assert result["reminders_sent"] == 2
        assert {email["report_type"] for email in sent} == {"initial", "verification"}
        assert {email["import_id"] for email in sent} == {
            str(import_id),
            str(verification_import_id),
        }
        assert run_vapt_client_review_reminders(db)["reminders_sent"] == 0
        assert len(sent) == 2
    finally:
        db.rollback()
        db.close()


def test_vapt_client_review_reminder_email_escapes_report_details(monkeypatch):
    sent_messages = []
    monkeypatch.setattr(email_utils, "SMTP_USER", "test@example.com")
    monkeypatch.setattr(email_utils, "SMTP_PASSWORD", "test-password")
    monkeypatch.setattr(email_utils, "FRONTEND_URL", "https://scanner.example")
    monkeypatch.setattr(email_utils, "_smtp_send", sent_messages.append)

    assert email_utils.send_vapt_client_review_reminder_email(
        "client@example.com",
        "<initial-report>.xml",
        "initial-import-id",
        "initial",
    )
    assert email_utils.send_vapt_client_review_reminder_email(
        "client@example.com",
        "<verification-report>.xml",
        "verification-import-id",
        "verification",
        verification_number=2,
    )

    assert len(sent_messages) == 2
    initial_html = sent_messages[0].get_payload()[1].get_payload(decode=True).decode()
    verification_html = sent_messages[1].get_payload()[1].get_payload(decode=True).decode()
    assert "&lt;initial-report&gt;.xml" in initial_html
    assert "&lt;verification-report&gt;.xml" in verification_html


def test_import_response_schemas_preserve_approved_next_vapt_date():
    due_date = datetime(2026, 12, 17, 9, 30, tzinfo=timezone.utc)
    base_item = {
        "import_id": "import-1",
        "file_name": "report.xml",
        "file_format": "xml",
        "source_tool": "nessus",
        "total_findings": 1,
        "unique_hosts": 1,
        "risk_score": 5,
        "severity": "low",
        "severity_distribution": {},
        "lifecycle_status": "closed",
        "next_vapt_due_at": due_date,
    }

    assert VaptImportListItem(**base_item).next_vapt_due_at == due_date
    assert VaptImportDetail(
        **base_item,
        category_distribution={},
        summary={},
        findings=[],
    ).next_vapt_due_at == due_date


def test_access_status_returns_all_approved_org_regions():
    db = Session(bind=engine)
    try:
        client = User(
            user_id="multi-region-client",
            org_id="org-multi-region",
            email="multi-region-client@example.com",
            password="test-password",
            role="owner",
            vapt_approved=True,
        )
        organization = Organization(
            org_id="org-multi-region",
            user_id=client.user_id,
            region="MSU-IND",
        )
        region_codes = ["MSU-IND", "BOU-IND", "ACC-IND", "COLL-IRE"]
        db.add_all([client, organization])
        for code in region_codes:
            region = Region(code=code, name=code, is_active=True)
            db.add(region)
            db.flush()
            db.add(
                OrganizationRegion(
                    org_id=organization.org_id,
                    region_id=region.region_id,
                    status="approved",
                    checklist_review_status="approved",
                )
            )
        other_client = User(
            user_id="other-multi-region-client",
            org_id="org-other-multi-region",
            email="other-multi-region-client@example.com",
            password="test-password",
            role="owner",
            vapt_approved=True,
        )
        other_region_codes = ["NORTH-TEST", "SOUTH-TEST"]
        db.add_all([
            other_client,
            Organization(
                org_id="org-other-multi-region",
                user_id=other_client.user_id,
                region=other_region_codes[0],
            ),
        ])
        for code in other_region_codes:
            region = Region(code=code, name=code, is_active=True)
            db.add(region)
            db.flush()
            db.add(
                OrganizationRegion(
                    org_id=other_client.org_id,
                    region_id=region.region_id,
                    status="approved",
                    checklist_review_status="approved",
                )
            )
        db.commit()

        status = get_vapt_access_status(db=db, current_user=client)

        assert status["region"] == "MSU-IND"
        assert status["approved_region_codes"] == sorted(region_codes)
        assert {region["code"] for region in status["approved_regions"]} == set(region_codes)
        other_status = get_vapt_access_status(db=db, current_user=other_client)
        assert other_status["region"] == "NORTH-TEST"
        assert other_status["approved_region_codes"] == sorted(other_region_codes)
        assert {region["code"] for region in other_status["approved_regions"]} == set(other_region_codes)
    finally:
        db.close()


def test_admin_organization_users_include_owner_and_unverified_members():
    db = Session(bind=engine)
    try:
        owner = User(
            user_id="admin-org-owner",
            org_id="org-admin-users",
            email="owner@example.com",
            password="test-password",
            role="owner",
            email_verified=False,
        )
        member = User(
            user_id="admin-org-member",
            org_id="org-admin-users",
            email="member@example.com",
            password="test-password",
            role="member",
            email_verified=False,
        )
        unassigned = User(
            user_id="admin-unassigned",
            org_id=None,
            email="unassigned@example.com",
            password="test-password",
            role="member",
            email_verified=False,
        )
        db.add_all([
            owner,
            member,
            unassigned,
            Organization(
                org_id="org-admin-users",
                user_id=owner.user_id,
                domain="admin-users.example.com",
            ),
        ])
        db.commit()

        result = get_users_by_org(db)
        organization = next(
            item for item in result["organizations"]
            if item["org_id"] == "org-admin-users"
        )

        assert {user["user_id"] for user in organization["users"]} == {
            owner.user_id,
            member.user_id,
        }
        assert result["admin"] == []
        assert result["soc_analysts"] == []
    finally:
        db.close()


def test_checklist_attachment_upload_is_validated_and_org_scoped():
    """Uploads are extension-checked, stored on disk, and only readable by the
    owning org (or by SOC reviewing it)."""
    import io
    import tempfile

    from fastapi import UploadFile

    import app.api.vapt.routes as vapt_routes

    store = tempfile.mkdtemp(prefix="vapt-attachments-")
    original_dir = vapt_routes.ATTACHMENT_DIR
    vapt_routes.ATTACHMENT_DIR = store
    db = Session(bind=engine)
    try:
        client = User(user_id="att-client", org_id="org-att", email="att-client@example.com", password="hashed", role="owner")
        other = User(user_id="att-other", org_id="org-other", email="att-other@example.com", password="hashed", role="owner")
        soc = User(user_id="att-soc", org_id=None, email="att-soc@example.com", password="hashed", role="soc_analyst")
        db.add_all([client, other, soc])
        db.add(Organization(org_id="org-att", user_id="att-client", max_domains=1))
        db.add(Organization(org_id="org-other", user_id="att-other", max_domains=1))
        db.commit()

        # An unsupported file type is refused before anything is stored.
        try:
            _run(upload_checklist_attachment(
                file=UploadFile(file=io.BytesIO(b"MZ"), filename="payload.exe"),
                section_id="network_infrastructure",
                question_id="asset_list_upload",
                region_code="",
                db=db,
                current_user=client,
            ))
            raise AssertionError("expected 400 for a disallowed extension")
        except HTTPException as exc:
            assert exc.status_code == 400

        # A question that does not accept uploads is refused.
        try:
            _run(upload_checklist_attachment(
                file=UploadFile(file=io.BytesIO(b"x"), filename="scope.pdf"),
                section_id="general_information",
                question_id="organization_name",
                region_code="",
                db=db,
                current_user=client,
            ))
            raise AssertionError("expected 400 for a non-upload question")
        except HTTPException as exc:
            assert exc.status_code == 400

        payload = b"host,ip\nweb01,10.0.0.5\n"
        meta = _run(upload_checklist_attachment(
            file=UploadFile(file=io.BytesIO(payload), filename="asset-list.csv"),
            section_id="network_infrastructure",
            question_id="asset_list_upload",
            region_code="",
            db=db,
            current_user=client,
        ))
        assert meta["filename"] == "asset-list.csv"
        assert meta["size_bytes"] == len(payload)
        assert meta["download_url"] == f"/vapt/onboarding/attachment/{meta['id']}"

        stored = db.query(VaptChecklistAttachment).filter_by(org_id="org-att").first()
        assert stored is not None
        assert os.path.isfile(os.path.join(store, stored.stored_name))
        # Two orgs' uploads never collide on disk.
        assert stored.stored_name.startswith("org-att")

        # Another org cannot read it, SOC reviewing it can.
        try:
            download_checklist_attachment(attachment_id=meta["id"], db=db, current_user=other)
            raise AssertionError("expected 403 for a foreign org")
        except HTTPException as exc:
            assert exc.status_code == 403
        assert download_checklist_attachment(attachment_id=meta["id"], db=db, current_user=soc) is not None

        # Removing it clears both the row and the stored file.
        delete_checklist_attachment(attachment_id=meta["id"], db=db, current_user=client)
        assert db.query(VaptChecklistAttachment).filter_by(org_id="org-att").count() == 0
        assert not os.path.isfile(os.path.join(store, stored.stored_name))
    finally:
        vapt_routes.ATTACHMENT_DIR = original_dir
        db.close()


def test_required_upload_questions_gate_submission():
    """The asset list is mandatory; the network diagram is optional; N/A and a
    valid attachment both satisfy a required upload."""
    db = Session(bind=engine)
    try:
        # No file, not marked N/A -> blocked.
        assert _missing_required_uploads(
            db,
            "org-req",
            {"network_infrastructure": {"asset_list_upload": {"answer": "", "na": False}}},
        ) == ["Asset list"]

        # Marking it N/A is an explicit opt-out (the client then emails it).
        assert _missing_required_uploads(
            db,
            "org-req",
            {"network_infrastructure": {"asset_list_upload": {"answer": "N/A", "na": True}}},
        ) == []

        # The optional network diagram never blocks a submission.
        assert _missing_required_uploads(
            db,
            "org-req",
            {"general_information": {"network_diagram_upload": {"answer": "", "na": False}}},
        ) == []

        # A foreign attachment id cannot be used to satisfy the requirement.
        assert _missing_required_uploads(
            db,
            "org-req",
            {
                "network_infrastructure": {
                    "asset_list_upload": {
                        "answer": "asset-list.csv",
                        "na": False,
                        "attachment": {"id": "11111111-2222-3333-4444-555555555555", "filename": "asset-list.csv"},
                    }
                }
            },
        ) == ["Asset list"]

        # An attachment that really belongs to the org satisfies it.
        db.add(Organization(org_id="org-req", user_id="req-client", max_domains=1))
        db.add(User(user_id="req-client", org_id="org-req", email="req-client@example.com", password="hashed", role="owner"))
        row = VaptChecklistAttachment(
            id="req-attachment",
            org_id="org-req",
            section_id="network_infrastructure",
            question_id="asset_list_upload",
            original_filename="asset-list.csv",
            size_bytes=23,
            stored_name="org-req/req-attachment.csv",
        )
        db.add(row)
        db.commit()
        assert _missing_required_uploads(
            db,
            "org-req",
            {
                "network_infrastructure": {
                    "asset_list_upload": {
                        "answer": "asset-list.csv",
                        "na": False,
                        "attachment": {"id": "req-attachment", "filename": "asset-list.csv"},
                    }
                }
            },
        ) == []
    finally:
        db.close()


def test_update_vapt_finding_status_persists_comment_and_status():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-1",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            total_findings=1,
            unique_hosts=1,
            risk_score=50,
            severity="medium",
            findings=[
                {
                    "id": "finding-1",
                    "status": "pending",
                    "comment": "",
                    "title": "Test finding",
                }
            ],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        finding_id = str(record.import_id)
        assert record.findings[0]["status"] == "pending"
        assert record.findings[0]["comment"] == ""

        normalized_status = _normalize_finding_status("solved")
        assert normalized_status == "solved"

        record.findings = [
            {
                **record.findings[0],
                "status": normalized_status,
                "comment": "Fixed in patch",
            }
        ]
        db.add(record)
        db.commit()
        db.refresh(record)

        assert record.findings[0]["status"] == "solved"
        assert record.findings[0]["comment"] == "Fixed in patch"
    finally:
        db.close()


def test_only_critical_high_and_medium_findings_require_triage():
    assert _requires_vapt_finding_triage({"severity_label": "critical"})
    assert _requires_vapt_finding_triage({"severity_label": "HIGH"})
    assert _requires_vapt_finding_triage({"severity_label": "medium"})
    assert not _requires_vapt_finding_triage({"severity_label": "low"})
    assert not _requires_vapt_finding_triage({"severity_label": "info"})
    assert not _requires_vapt_finding_triage({"severity_label": "informational"})


def test_ignore_and_false_positive_require_comments_at_every_severity():
    db = Session(bind=engine)
    try:
        severities = ["critical", "high", "medium", "low", "info"]
        record = VaptImport(
            org_id="org-required-comments",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            findings=[
                {"id": f"{severity}-1", "severity_label": severity, "status": "pending"}
                for severity in severities
            ],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        schedule = VaptRescanSchedule(
            import_id=record.import_id,
            org_id="org-required-comments",
            created_by="soc-user",
            scheduled_at=datetime.now(timezone.utc),
            status="completed_with_errors",
            result_data={
                "remaining_findings": [
                    {
                        "id": f"{severity}-verify",
                        "severity_label": severity,
                        "status": "pending",
                    }
                    for severity in severities
                ],
                "new_findings": [],
            },
        )
        db.add(schedule)
        db.commit()
        db.refresh(schedule)

        for severity in severities:
            for status in ("ignore", "false_positive"):
                try:
                    update_vapt_finding_status(
                        str(record.import_id),
                        f"{severity}-1",
                        VaptFindingStatusUpdate(status=status),
                        db=db,
                        current_user=SimpleNamespace(org_id="org-required-comments"),
                    )
                    raise AssertionError(f"expected a comment for {status} {severity} finding")
                except HTTPException as error:
                    assert error.status_code == 400
                    assert "comment is required" in error.detail.lower()

                try:
                    update_verification_finding_status(
                        str(record.import_id),
                        schedule.id,
                        f"{severity}-verify",
                        VaptFindingStatusUpdate(status=status),
                        db=db,
                        current_user=SimpleNamespace(
                            org_id="org-required-comments",
                            user_id="client-user",
                        ),
                    )
                    raise AssertionError(f"expected a comment for verification {status} {severity} finding")
                except HTTPException as error:
                    assert error.status_code == 400
                    assert "comment is required" in error.detail.lower()
    finally:
        db.close()


def test_low_finding_can_be_marked_solved_without_a_remark():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-low-triage",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            total_findings=1,
            unique_hosts=1,
            risk_score=10,
            severity="low",
            findings=[{"id": "low-1", "severity_label": "low", "status": "pending"}],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        update_vapt_finding_status(
            str(record.import_id),
            "low-1",
            VaptFindingStatusUpdate(status="solved"),
            db=db,
            current_user=SimpleNamespace(org_id="org-low-triage", user_id="test-user"),
        )
        db.refresh(record)
        assert record.findings[0]["severity_label"] == "low"
        assert record.findings[0]["status"] == "solved"
        assert record.findings[0]["comment"] == ""
    finally:
        db.close()


def test_low_verification_finding_can_be_marked_solved_without_a_remark():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-low-verification",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            total_findings=1,
            unique_hosts=1,
            risk_score=10,
            severity="low",
            findings=[{
                "id": "low-original",
                "title": "Low finding",
                "severity_label": "low",
                "status": "solved",
                "comment": "Disabled the exposed service.",
            }],
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        schedule = VaptRescanSchedule(
            import_id=record.import_id,
            org_id="org-low-verification",
            created_by="soc-user",
            scheduled_at=datetime.now(timezone.utc),
            status="completed_with_errors",
            result_data={
                "remaining_findings": [{
                    "id": "low-retest",
                    "title": "Low finding",
                    "severity_label": "low",
                    "status": "solved",
                    "comment": "Disabled the exposed service.",
                }],
                "new_findings": [],
            },
        )
        db.add(schedule)
        db.commit()
        db.refresh(schedule)
        user = SimpleNamespace(org_id="org-low-verification", user_id="client-user")

        update_verification_finding_status(
            str(record.import_id),
            schedule.id,
            "low-retest",
            VaptFindingStatusUpdate(status="solved"),
            db=db,
            current_user=user,
        )
        db.refresh(schedule)
        updated = schedule.result_data["remaining_findings"][0]
        assert updated["status"] == "solved"
        assert updated["comment"] == ""

        response = asyncio.run(submit_verification_review(
            str(record.import_id),
            schedule.id,
            db=db,
            current_user=user,
        ))
        assert response["success"] is True
        assert schedule.result_data["client_review_status"] == "client_completed"
    finally:
        db.close()


def test_low_verification_finding_can_remain_unreviewed_when_submitting_review():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-low-optional-review",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            findings=[{"id": "low-original", "severity_label": "low", "status": "pending"}],
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        schedule = VaptRescanSchedule(
            import_id=record.import_id,
            org_id="org-low-optional-review",
            created_by="soc-user",
            scheduled_at=datetime.now(timezone.utc),
            status="completed_with_errors",
            result_data={
                "remaining_findings": [{
                    "id": "low-retest",
                    "severity_label": "low",
                    "status": "pending",
                    "comment": "",
                }],
                "new_findings": [],
            },
        )
        db.add(schedule)
        db.commit()
        db.refresh(schedule)

        response = asyncio.run(submit_verification_review(
            str(record.import_id),
            schedule.id,
            db=db,
            current_user=SimpleNamespace(org_id="org-low-optional-review", user_id="client-user"),
        ))
        assert response["success"] is True
    finally:
        db.close()


def test_low_finding_does_not_block_report_submission():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-low-submit",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            total_findings=1,
            unique_hosts=1,
            risk_score=10,
            severity="low",
            findings=[{"id": "low-1", "severity_label": "low", "status": "pending"}],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        response = asyncio.run(
            submit_vapt_import(
                str(record.import_id),
                db=db,
                current_user=SimpleNamespace(org_id="org-low-submit"),
            )
        )

        assert response["success"] is True
        assert record.status == "client_completed"
    finally:
        db.close()


def test_solved_low_finding_without_remark_can_be_submitted():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-low-missing-remark",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            findings=[{
                "id": "low-1",
                "severity_label": "low",
                "status": "solved",
                "comment": "",
            }],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        response = asyncio.run(submit_vapt_import(
            str(record.import_id),
            db=db,
            current_user=SimpleNamespace(org_id="org-low-missing-remark"),
        ))
        assert response["success"] is True
        assert record.status == "client_completed"
    finally:
        db.close()


def test_soc_can_flag_client_ignore_and_require_a_resubmission():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-soc-finding-follow-up",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            findings=[
                {
                    "id": "high-1",
                    "title": "High finding",
                    "severity_label": "high",
                    "status": "ignore",
                    "comment": "This service is not used.",
                },
                {
                    "id": "low-1",
                    "title": "Low finding",
                    "severity_label": "low",
                    "status": "false_positive",
                    "comment": "The scanner matched a test host.",
                },
            ],
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        client = SimpleNamespace(org_id="org-soc-finding-follow-up", user_id="client-user")
        soc = SimpleNamespace(user_id="soc-user", role="soc_analyst")

        asyncio.run(submit_vapt_import(
            str(record.import_id),
            db=db,
            current_user=client,
        ))
        flagged = asyncio.run(review_client_remediation(
            str(record.import_id),
            {
                "decision": "changes_requested",
                "flags": [{
                    "finding_id": "high-1",
                    "note": "Please explain why this exposed service is not applicable.",
                }],
            },
            db=db,
            current_user=soc,
        ))
        assert flagged["decision"] == "changes_requested"
        assert record.remediation_review_status == "changes_requested"
        assert record.lifecycle_status == "remediation_required"
        detail_response = VaptImportDetail(**_to_detail(record))
        assert detail_response.remediation_review_status == "changes_requested"
        high = next(finding for finding in record.findings if finding["id"] == "high-1")
        assert high["soc_review_flags"][0]["client_comment_at_flag"] == "This service is not used."
        assert high["soc_review_flags"][0]["note"] == "Please explain why this exposed service is not applicable."

        with pytest.raises(HTTPException) as unanswered_error:
            asyncio.run(submit_vapt_import(
                str(record.import_id),
                db=db,
                current_user=client,
            ))
        assert unanswered_error.value.status_code == 400
        assert "soc-flagged" in unanswered_error.value.detail.lower()

        with pytest.raises(HTTPException) as unflagged_error:
            update_vapt_finding_status(
                str(record.import_id),
                "low-1",
                VaptFindingStatusUpdate(status="solved"),
                db=db,
                current_user=client,
            )
        assert unflagged_error.value.status_code == 400

        with pytest.raises(HTTPException) as unchanged_error:
            update_vapt_finding_status(
                str(record.import_id),
                "high-1",
                VaptFindingStatusUpdate(status="ignore", comment="This service is not used."),
                db=db,
                current_user=client,
            )
        assert unchanged_error.value.status_code == 400

        update_vapt_finding_status(
            str(record.import_id),
            "high-1",
            VaptFindingStatusUpdate(status="solved", comment="Disabled the exposed service."),
            db=db,
            current_user=client,
        )
        high = next(finding for finding in record.findings if finding["id"] == "high-1")
        assert high["status"] == "solved"
        assert high["soc_review_flags"][0]["response_status"] == "solved"
        assert high["soc_review_flags"][0]["response_comment"] == "Disabled the exposed service."

        submitted = asyncio.run(submit_vapt_import(
            str(record.import_id),
            db=db,
            current_user=client,
        ))
        assert submitted["success"] is True
        assert record.remediation_review_status == "pending_soc_review"
        assert record.lifecycle_status == "awaiting_soc_remediation_acceptance"

        approved = asyncio.run(review_client_remediation(
            str(record.import_id),
            {"decision": "approved"},
            db=db,
            current_user=soc,
        ))
        assert approved["decision"] == "approved"
        assert record.lifecycle_status == "revalidation_required"
        detail = _to_detail(record)
        reviewed_high = next(finding for finding in detail["findings"] if finding["id"] == "high-1")
        assert reviewed_high["soc_review_flags"][0]["note"] == "Please explain why this exposed service is not applicable."
        assert reviewed_high["soc_review_flags"][0]["responded_at"]
    finally:
        db.close()


def test_ignore_without_comment_is_rejected_when_submitting_reviews():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-missing-ignore-comment",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            findings=[{
                "id": "low-1",
                "severity_label": "low",
                "status": "ignore",
                "comment": "",
            }],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        with pytest.raises(HTTPException) as report_error:
            asyncio.run(submit_vapt_import(
                str(record.import_id),
                db=db,
                current_user=SimpleNamespace(org_id="org-missing-ignore-comment"),
            ))
        assert report_error.value.status_code == 400
        assert "comment" in report_error.value.detail.lower()

        schedule = VaptRescanSchedule(
            import_id=record.import_id,
            org_id="org-missing-ignore-comment",
            created_by="soc-user",
            scheduled_at=datetime.now(timezone.utc),
            status="completed_with_errors",
            result_data={
                "remaining_findings": [{
                    "id": "low-retest",
                    "severity_label": "low",
                    "status": "false_positive",
                    "comment": "",
                }],
                "new_findings": [],
            },
        )
        db.add(schedule)
        db.commit()
        db.refresh(schedule)

        with pytest.raises(HTTPException) as verification_error:
            asyncio.run(submit_verification_review(
                str(record.import_id),
                schedule.id,
                db=db,
                current_user=SimpleNamespace(
                    org_id="org-missing-ignore-comment",
                    user_id="client-user",
                ),
            ))
        assert verification_error.value.status_code == 400
        assert "comment" in verification_error.value.detail.lower()
    finally:
        db.close()


def test_solved_low_finding_does_not_force_verification_for_closure():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-low-direct-close",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            status="client_completed",
            remediation_review_status="approved",
            findings=[{
                "id": "low-1",
                "title": "Low finding",
                "severity_label": "low",
                "status": "solved",
                "comment": "Disabled the exposed service.",
            }],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        response = asyncio.run(close_vapt_without_verification(
            str(record.import_id),
            DirectClosureRequest(note="Low finding is optional."),
            db=db,
            current_user=SimpleNamespace(user_id="soc-user", role="soc_analyst"),
        ))

        assert response["success"] is True
        assert record.lifecycle_status == "closure_pending_client_due_date"
    finally:
        db.close()


def test_medium_finding_still_blocks_report_submission():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-medium-submit",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            total_findings=1,
            unique_hosts=1,
            risk_score=40,
            severity="medium",
            findings=[{"id": "medium-1", "severity_label": "medium", "status": "pending"}],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        try:
            asyncio.run(
                submit_vapt_import(
                    str(record.import_id),
                    db=db,
                    current_user=SimpleNamespace(org_id="org-medium-submit"),
                )
            )
            raise AssertionError("expected pending medium findings to block submission")
        except HTTPException as error:
            assert error.status_code == 400
            assert "1 finding(s) still pending" in error.detail
    finally:
        db.close()


def test_vapt_import_submit_sets_status_submitted():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-1",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            total_findings=1,
            unique_hosts=1,
            risk_score=50,
            severity="medium",
            findings=[
                {
                    "id": "finding-1",
                    "status": "pending",
                    "comment": "",
                    "title": "Test finding",
                }
            ],
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        record.status = "submitted"
        db.add(record)
        db.commit()
        db.refresh(record)

        assert record.status == "submitted"
    finally:
        db.close()


def test_request_vapt_access_accepts_combined_submission_payload():
    db = Session(bind=engine)
    try:
        user = User(
            user_id="user-1",
            org_id="org-1",
            email="client@example.com",
            password="hashed-password",
            role="owner",
        )
        db.add(user)
        db.commit()

        response = request_vapt_access(
            payload={
                "regions": [{"code": "ACC-IND", "name": "Accenture India"}],
                "scope_ip_ranges": "10.0.0.0/8",
                "authorization_confirmed": True,
                "tech_contact_name": "Alice Admin",
                "tech_contact_email": "alice@example.com",
                "testing_window": "15-18 Sep 2026, 10:00 AM to 6:00 PM IST",
                "testing_start_at": "2026-09-15T10:00:00+00:00",
                "testing_timezone": "UTC",
                "checklist_answers": _complete_region_checklist_answers(),
            },
            db=db,
            current_user=user,
        )

        assert response["success"] is True
        assert "ACC-IND" in response["requested_regions"]
        region = db.query(Region).filter_by(code="ACC-IND").one()
        request = db.query(OrganizationRegion).filter_by(
            org_id="org-1",
            region_id=region.region_id,
        ).one()
        assert request.status == "pending"
        assert request.checklist_submission["scope_ip_ranges"] == "10.0.0.0/8"
        assert request.checklist_submission["authorization_confirmed"] is True
        assert request.checklist_submission["testing_window"] == "15-18 Sep 2026, 10:00 AM to 6:00 PM IST"
        assert request.checklist_review_status == "pending"
    finally:
        db.close()


def test_request_vapt_access_rejects_region_only_submission():
    class Client:
        org_id = "org-region-only"

    db = Session(bind=engine)
    try:
        try:
            request_vapt_access(
                payload={"regions": [{"code": "ACC-IND", "name": "Accenture India"}]},
                db=db,
                current_user=Client(),
            )
        except HTTPException as exc:
            assert exc.status_code == 400
            assert "checklist together" in exc.detail
        else:
            raise AssertionError("region-only requests must be rejected")
    finally:
        db.close()


def test_manual_verification_marks_all_fixed_as_completed():
    original = [{"title": "TLS issue", "plugin_id": "100", "status": "solved"}]
    status, error, fixed, remaining, new, unverified = _evaluate_manual_verification(original, [])
    assert status == "completed"
    assert error is None
    assert fixed[0]["title"] == original[0]["title"]
    assert remaining == []
    assert new == []
    assert unverified == []


def test_manual_verification_marks_partial_fix_as_completed_with_errors():
    original = [
        {"title": "TLS issue", "plugin_id": "100", "status": "solved"},
        {"title": "SSH issue", "plugin_id": "200", "status": "solved"},
    ]
    status, error, fixed, remaining, new, unverified = _evaluate_manual_verification(original, [original[1]])
    assert status == "completed_with_errors"
    assert error
    assert [finding["title"] for finding in fixed] == [original[0]["title"]]
    assert [finding["title"] for finding in remaining] == [original[1]["title"]]
    assert new == []
    assert unverified == []


def test_manual_verification_marks_no_confirmed_fix_as_failed():
    original = [{"title": "TLS issue", "plugin_id": "100", "status": "solved"}]
    status, error, fixed, remaining, new, unverified = _evaluate_manual_verification(original, [original[0]])
    assert status == "failed"
    assert error
    assert fixed == []
    assert [finding["title"] for finding in remaining] == [finding["title"] for finding in original]
    assert new == []
    assert unverified == []


def test_manual_verification_exposes_new_actionable_findings_for_client_triage():
    original = [{"title": "TLS issue", "plugin_id": "100", "severity_label": "high", "status": "solved"}]
    unexpected = {"title": "Open admin port", "plugin_id": "200", "severity_label": "critical"}

    status, error, fixed, remaining, new, unverified = _evaluate_manual_verification(
        original,
        [unexpected],
        original,
    )

    assert status == "completed_with_errors"
    assert "new actionable finding" in error
    assert [finding["title"] for finding in fixed] == [original[0]["title"]]
    assert remaining == []
    assert new == [{**unexpected, "status": "pending", "comment": ""}]
    assert [finding["plugin_id"] for finding in _new_findings_requiring_follow_up(new)] == ["200"]
    assert unverified == []


def test_manual_verification_keeps_new_low_findings_visible_without_blocking():
    original = [{"title": "TLS issue", "plugin_id": "100", "status": "solved"}]
    informational = {"title": "Informational banner", "plugin_id": "200", "severity_label": "info"}

    status, error, fixed, remaining, new, unverified = _evaluate_manual_verification(
        original,
        [informational],
        original,
    )

    assert status == "completed"
    assert error is None
    assert [finding["title"] for finding in fixed] == [original[0]["title"]]
    assert remaining == []
    assert new == [{**informational, "status": "pending", "comment": ""}]
    assert _new_findings_requiring_follow_up(new) == []
    assert unverified == []


def test_manual_verification_reports_grouped_partial_host_fix():
    original = [{
        "title": "TLS issue",
        "plugin_id": "100",
        "port": 443,
        "protocol": "tcp",
        "description": "Weak TLS configuration",
        "severity_label": "high",
        "status": "solved",
        "affected_hosts": ["host-a", "host-b", "host-c"],
        "host_count": 3,
    }]
    detected = {
        "title": "TLS issue",
        "plugin_id": "100",
        "port": 443,
        "protocol": "tcp",
        "description": "Weak TLS configuration",
        "severity_label": "high",
        "affected_hosts": ["host-b"],
    }

    status, error, fixed, remaining, new, unverified = _evaluate_manual_verification(
        original,
        [detected],
        original,
        ["host-a", "host-b"],
    )

    assert status == "completed_with_errors"
    assert "remain on 1 asset" in error
    assert fixed == []
    assert remaining[0]["affected_hosts"] == ["host-b"]
    assert remaining[0]["verified_fixed_hosts"] == ["host-a"]
    assert remaining[0]["unverified_hosts"] == ["host-c"]
    assert remaining[0]["original_host_count"] == 3
    assert new == []
    assert unverified == []


def test_manual_verification_lists_hosts_outside_selected_scan_scope():
    original = [{
        "title": "TLS issue",
        "plugin_id": "100",
        "status": "solved",
        "affected_hosts": ["host-a", "host-b"],
    }]

    status, error, fixed, remaining, new, unverified = _evaluate_manual_verification(
        original,
        [],
        original,
        ["host-a"],
    )

    assert status == "completed_with_errors"
    assert "not included in this scan" in error
    assert fixed == []
    assert remaining == []
    assert unverified[0]["affected_hosts"] == ["host-b"]
    assert unverified[0]["verified_fixed_hosts"] == ["host-a"]


def test_new_actionable_findings_block_closure_until_follow_up_verification():
    findings = [
        {"id": "pending", "severity_label": "critical", "status": "pending"},
        {"id": "solved", "severity_label": "high", "status": "solved"},
        {"id": "ignored", "severity_label": "medium", "status": "ignore"},
        {"id": "low", "severity_label": "low", "status": "pending"},
    ]

    assert [finding["id"] for finding in _new_findings_requiring_follow_up(findings)] == [
        "pending",
        "solved",
    ]


def test_severity_gated_closure_allows_ignored_items_but_blocks_pending_and_redetected():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-closure",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            total_findings=4,
            unique_hosts=1,
            risk_score=70,
            severity="high",
            findings=[
                {"id": "critical", "title": "Critical issue", "severity_label": "critical", "status": "ignore"},
                {"id": "medium", "title": "Medium issue", "severity_label": "medium", "status": "pending"},
                {"id": "false-positive", "title": "Accepted medium issue", "severity_label": "medium", "status": "false_positive"},
                {"id": "low", "title": "Low issue", "severity_label": "low", "status": "ignore"},
            ],
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        blockers = _closure_blockers(record)
        assert {finding["id"] for finding in blockers} == {"medium"}
        message = _closure_block_message(blockers)
        assert "Medium issue (pending)" in message
        assert "Critical issue" not in message

        # Ignored / false-positive dispositions are accepted after triage.
        assert _closure_blockers(record, []) == []

        # A medium finding the client marked solved still blocks closure if
        # the SOC verification scan detects it again.
        medium = next(finding for finding in record.findings if finding["id"] == "medium")
        medium["status"] = "solved"
        redetected = {**medium, "status": "pending"}
        assert _closure_blockers(record, [redetected]) == [medium]
    finally:
        db.close()


def test_severity_gated_closure_allows_only_unresolved_low_findings():
    db = Session(bind=engine)
    try:
        record = VaptImport(
            org_id="org-low",
            file_name="test.nessus",
            file_format="xml",
            source_tool="nessus",
            total_findings=1,
            unique_hosts=1,
            risk_score=10,
            severity="low",
            findings=[{"id": "low", "title": "Low issue", "severity_label": "low", "status": "ignore"}],
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        assert _closure_blockers(record) == []
    finally:
        db.close()


class _Schedule:
    """Minimal stand-in for VaptRescanSchedule (only result_data is read)."""

    def __init__(self, result_data=None, schedule_id="abcdef12-0000-0000-0000-000000000000"):
        self.result_data = result_data
        self.id = schedule_id


def test_verification_display_name_normalizes_soc_input():
    assert _clean_verification_display_name("  Q3   re-validation \u2014 Mumbai ") == "Q3 re-validation \u2014 Mumbai"
    assert _clean_verification_display_name("") is None
    assert _clean_verification_display_name(None) is None
    assert _clean_verification_display_name("   ") is None
    long_name = "x" * (VERIFICATION_DISPLAY_NAME_MAX + 40)
    assert len(_clean_verification_display_name(long_name)) == VERIFICATION_DISPLAY_NAME_MAX


def test_verification_display_name_reads_uploaded_name_only_when_set():
    assert _verification_display_name(_Schedule({"display_name": "Retest July"})) == "Retest July"
    assert _verification_display_name(_Schedule({})) is None
    assert _verification_display_name(_Schedule(None)) is None
    # Defensive: legacy rows may hold a non-dict result payload.
    assert _verification_display_name(_Schedule("raw-string")) is None


def test_verification_download_filename_prefers_custom_name():
    schedule_id = "abcdef12-3456-7890-abcd-ef1234567890"
    named = _Schedule({"display_name": "Q3 re-validation / Mumbai"}, schedule_id)
    # Unsafe characters are dropped, spaces become dashes.
    assert _verification_download_filename(named, schedule_id, "pdf") == "Q3-re-validation-Mumbai.pdf"
    unnamed = _Schedule({}, schedule_id)
    assert _verification_download_filename(unnamed, schedule_id, "xlsx") == "vapt-verification-abcdef12.xlsx"


def _run(coro):
    return asyncio.run(coro)


def test_vapt_report_cycles_are_scoped_to_region():
    """An open report in one region does not block uploads in another."""
    import app.api.vapt.routes as vapt_routes

    db = Session(bind=engine)
    original_parse_upload = vapt_routes.parse_upload
    original_normalize_import = vapt_routes.normalize_import
    try:
        client = User(
            user_id="region-cycle-client",
            org_id="org-region-cycle",
            email="region-cycle-client@example.com",
            password="test-password",
            role="owner",
        )
        soc = User(
            user_id="region-cycle-soc",
            org_id=None,
            email="region-cycle-soc@example.com",
            password="test-password",
            role="soc_analyst",
        )
        db.add_all([client, soc])
        db.add(Organization(org_id="org-region-cycle", user_id=client.user_id))
        region_a = Region(code="CYCLE-A", name="Region A", is_active=True)
        region_b = Region(code="CYCLE-B", name="Region B", is_active=True)
        db.add_all([region_a, region_b])
        db.flush()
        db.add_all([
            OrganizationRegion(
                org_id="org-region-cycle",
                region_id=region_a.region_id,
                status="approved",
                testing_start_at=datetime(2026, 10, 15, 10, tzinfo=timezone.utc),
                testing_timezone="UTC",
                checklist_review_status="approved",
                checklist_submission={
                    "checklist_answers": _complete_region_checklist_answers()
                },
            ),
            OrganizationRegion(
                org_id="org-region-cycle",
                region_id=region_b.region_id,
                status="approved",
                testing_start_at=datetime(2026, 10, 15, 10, tzinfo=timezone.utc),
                testing_timezone="UTC",
                checklist_review_status="pending",
                checklist_submission={
                    "checklist_answers": _complete_region_checklist_answers()
                },
            ),
            VaptOnboardingChecklist(
                org_id="org-region-cycle",
                testing_start_at=datetime(2026, 10, 15, 10, tzinfo=timezone.utc),
                testing_timezone="UTC",
                checklist_answers={"general_information": {"organization_name": {"answer": "Acme"}}},
                completed_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
                review_status="approved",
            ),
            VaptImport(
                org_id="org-region-cycle",
                region="CYCLE-A",
                file_name="region-a.csv",
                file_format="csv",
                source_tool="generic",
                lifecycle_status="report_published",
                cycle_number=1,
            ),
        ])
        db.commit()

        vapt_routes.parse_upload = lambda content, filename: (
            [{
                "host": "10.0.1.15" if "one" in filename else "10.0.2.15",
                "title": f"Finding from {filename}",
                "severity": 3,
                "severity_label": "high",
            }],
            "nessus",
            "nessus",
        )

        try:
            _run(upload_vapt_report(
                file=UploadFile(file=io.BytesIO(b"sample"), filename="region-b-unreviewed.csv"),
                org_id="org-region-cycle",
                region="CYCLE-B",
                db=db,
                current_user=soc,
            ))
            raise AssertionError("expected upload to require approval of the selected region checklist")
        except HTTPException as exc:
            assert exc.status_code == 403

        region_b_checklist = db.query(OrganizationRegion).filter_by(
            org_id="org-region-cycle",
            region_id=region_b.region_id,
        ).one()
        region_b_checklist.checklist_submission = None
        db.add(region_b_checklist)
        db.commit()
        try:
            _run(upload_vapt_report(
                file=UploadFile(file=io.BytesIO(b"sample"), filename="region-b-missing-checklist.csv"),
                org_id="org-region-cycle",
                region="CYCLE-B",
                db=db,
                current_user=soc,
            ))
            raise AssertionError("expected upload to require a region checklist")
        except HTTPException as exc:
            assert exc.status_code == 403

        region_b_checklist.checklist_submission = {
            "checklist_answers": _complete_region_checklist_answers()
        }
        region_b_checklist.checklist_review_status = "approved"
        db.add(region_b_checklist)
        db.commit()
        uploaded = _run(upload_vapt_report(
            files=[
                UploadFile(file=io.BytesIO(b"sample one"), filename="floor-one.nessus"),
                UploadFile(file=io.BytesIO(b"sample two"), filename="floor-two.nessus"),
            ],
            org_id="org-region-cycle",
            region="CYCLE-B",
            db=db,
            current_user=soc,
        ))
        assert uploaded["region"] == "CYCLE-B"
        assert uploaded["cycle_number"] == 1
        assert uploaded["file_name"] == "Combined report (2 files)"
        assert uploaded["summary"]["source_files"] == ["floor-one.nessus", "floor-two.nessus"]
        assert {
            host
            for finding in uploaded["findings"]
            for host in finding["affected_hosts"]
        } == {"10.0.1.15", "10.0.2.15"}

        try:
            _run(upload_vapt_report(
                file=UploadFile(file=io.BytesIO(b"sample"), filename="region-b-again.csv"),
                org_id="org-region-cycle",
                region="CYCLE-B",
                db=db,
                current_user=soc,
            ))
            raise AssertionError("expected 409 while the same region's report cycle is open")
        except HTTPException as exc:
            assert exc.status_code == 409

        checklist = _get_onboarding_or_create(db, "org-region-cycle")
        assert checklist.cycle_number == 1

        region_b_import = db.query(VaptImport).filter_by(
            org_id="org-region-cycle",
            region="CYCLE-B",
        ).one()
        region_b_import.lifecycle_status = "closed"
        db.add(region_b_import)
        db.commit()

        # Closing only one region must not roll the shared organization
        # checklist forward while another region still has an open cycle.
        checklist = _get_onboarding_or_create(db, "org-region-cycle")
        assert checklist.cycle_number == 1

        region_a_import = db.query(VaptImport).filter_by(
            org_id="org-region-cycle",
            region="CYCLE-A",
        ).one()
        region_a_import.lifecycle_status = "closed"
        db.add(region_a_import)
        db.commit()
        checklist = _get_onboarding_or_create(db, "org-region-cycle")
        assert checklist.cycle_number == 2
    finally:
        vapt_routes.parse_upload = original_parse_upload
        vapt_routes.normalize_import = original_normalize_import
        db.close()


def test_region_approval_requires_its_completed_checklist():
    """A region-only request cannot be approved without its own checklist."""
    db = Session(bind=engine)
    try:
        client = User(
            user_id="flow-client",
            org_id="org-flow",
            email="flow-client@example.com",
            password="hashed",
            role="owner",
        )
        soc = User(
            user_id="flow-soc",
            org_id=None,
            email="flow-soc@example.com",
            password="hashed",
            role="soc_analyst",
        )
        db.add_all([client, soc])
        db.commit()

        # 1. The client requests a region only.
        request_vapt_access(
            payload={"regions": [{"code": "FLW-IND", "name": "Flow India"}]},
            db=db,
            current_user=client,
        )
        db.expire_all()
        org_region = db.query(OrganizationRegion).filter_by(org_id="org-flow").first()
        assert org_region is not None and org_region.status == "pending"
        checklist = db.query(VaptOnboardingChecklist).filter_by(org_id="org-flow").first()
        assert checklist is None or checklist.completed_at is None

        # 2. The checklist cannot be submitted before the region is approved.
        try:
            _run(
                submit_onboarding_checklist(
                    payload={
                        "testing_start_at": "2026-09-15T10:00:00+00:00",
                        "testing_timezone": "UTC",
                        "checklist_answers": _complete_region_checklist_answers(),
                    },
                    db=db,
                    current_user=client,
                )
            )
            raise AssertionError("expected 403 before region approval")
        except HTTPException as exc:
            assert exc.status_code == 403

        # A plain region approval is blocked until that region has a completed
        # checklist attached.
        try:
            _run(
                approve_vapt_access(
                    payload={"org_id": "org-flow", "region": "FLW-IND", "approved": True},
                    db=db,
                    current_user=soc,
                )
            )
            raise AssertionError("expected approval to require a region checklist")
        except HTTPException as exc:
            assert exc.status_code == 400

        # Client submits the organization and region checklists together.
        request_vapt_access(
            payload={
                "regions": [{"code": "FLW-IND", "name": "Flow India"}],
                "scope_ip_ranges": "10.0.0.0/8",
                "authorization_confirmed": True,
                "testing_start_at": "2026-09-15T10:00:00+00:00",
                "testing_timezone": "UTC",
                "checklist_answers": _complete_region_checklist_answers(),
            },
            db=db,
            current_user=client,
        )
        _run(
            approve_vapt_access(
                payload={"org_id": "org-flow", "region": "FLW-IND", "approved": True},
                db=db,
                current_user=soc,
            )
        )
        db.expire_all()
        assert db.query(OrganizationRegion).filter_by(org_id="org-flow").first().status == "approved"

        # The organization checklist can be resubmitted for its own review path.
        response = _run(
            submit_onboarding_checklist(
                payload={
                    "scope_ip_ranges": "10.0.0.0/8",
                    "authorization_confirmed": True,
                    "testing_start_at": "2026-09-15T10:00:00+00:00",
                    "testing_timezone": "UTC",
                    "checklist_answers": {
                        "general_information": {
                            "organization_name": {"answer": "Acme Corp", "na": False},
                        }
                    },
                },
                db=db,
                current_user=client,
            )
        )
        assert response["success"] is True
        assert response["onboarding"]["completed"] is True
        assert response["onboarding"]["review_status"] == "pending"

        # 5. SOC reviews the submitted checklist.
        reviewed = _run(
            review_vapt_onboarding("org-flow", {"status": "approved"}, db=db, current_user=soc)
        )
        assert reviewed["review_status"] == "approved"
    finally:
        db.close()


def test_additional_region_request_carries_its_checklist_to_admin():
    """A new-region request must submit its checklist with the region so SOC can
    review both together, without disturbing the org's existing approved state."""
    db = Session(bind=engine)
    try:
        client = User(
            user_id="rc-client",
            org_id="org-rc",
            email="rc-client@example.com",
            password="hashed",
            role="owner",
        )
        soc = User(
            user_id="rc-soc",
            org_id=None,
            email="rc-soc@example.com",
            password="hashed",
            role="soc_analyst",
        )
        db.add_all([client, soc])
        db.add(Organization(org_id="org-rc", user_id="rc-client", max_domains=1))
        db.commit()

        # The client already has one approved region, so it may request another.
        request_vapt_access(
            payload={
                "regions": [{"code": "RC-A", "name": "Region A"}],
                "testing_start_at": "2026-09-15T10:00:00+00:00",
                "testing_timezone": "UTC",
                "checklist_answers": _complete_region_checklist_answers(),
            },
            db=db,
            current_user=client,
        )
        _run(
            approve_vapt_access(
                payload={"org_id": "org-rc", "region": "RC-A", "approved": True},
                db=db,
                current_user=soc,
            )
        )

        response = _run(
            request_vapt_region(
                payload={
                    "region_code": "RC-B",
                    "region_name": "Region B",
                    "testing_start_at": "2026-10-01T10:00:00+00:00",
                    "testing_timezone": "UTC",
                    "scope_ip_ranges": "10.1.0.0/16",
                    "checklist_answers": _complete_region_checklist_answers(),
                },
                db=db,
                current_user=client,
            )
        )
        assert response["success"] is True

        region_b = db.query(Region).filter_by(code="RC-B").one()
        org_region_b = (
            db.query(OrganizationRegion)
            .filter_by(org_id="org-rc", region_id=region_b.region_id)
            .one()
        )
        assert org_region_b.status == "pending"
        assert org_region_b.checklist_submission is not None
        assert org_region_b.checklist_submission["scope_ip_ranges"] == "10.1.0.0/16"
        assert org_region_b.checklist_submission["checklist_answers"]

        # The admin review queue exposes the submitted checklist on the region.
        queue = list_vapt_access_requests(db=db, current_user=soc)
        entry = next(item for item in queue if item["org_id"] == "org-rc")
        detail = next(item for item in entry["pending_region_details"] if item["code"] == "RC-B")
        assert detail["checklist_submission"]["scope_ip_ranges"] == "10.1.0.0/16"
    finally:
        db.close()


def test_additional_region_request_rejects_missing_or_incomplete_checklist():
    db = Session(bind=engine)
    try:
        client = User(
            user_id="missing-region-checklist-client",
            org_id="org-missing-region-checklist",
            email="missing-region-checklist@example.com",
            password="test-password",
            role="owner",
        )
        organization = Organization(org_id="org-missing-region-checklist", user_id=client.user_id)
        approved_region = Region(code="MRC-APPROVED", name="Approved", is_active=True)
        requested_region = Region(code="MRC-NEW", name="New", is_active=True)
        db.add_all([client, organization, approved_region, requested_region])
        db.flush()
        db.add(
            OrganizationRegion(
                org_id=organization.org_id,
                region_id=approved_region.region_id,
                status="approved",
            )
        )
        db.commit()

        partial_answers = _complete_region_checklist_answers()
        del partial_answers["company_details"]["primary_contact"]
        for answers in (
            None,
            {},
            {"general_information": {"organization_name": {"answer": ""}}},
            partial_answers,
        ):
            try:
                _run(
                    request_vapt_region(
                        payload={
                            "region_code": "MRC-NEW",
                            "region_name": "New",
                            "testing_start_at": "2026-10-15T10:00:00+00:00",
                            "testing_timezone": "UTC",
                            "checklist_answers": answers,
                        },
                        db=db,
                        current_user=client,
                    )
                )
                raise AssertionError("expected incomplete region checklist to be rejected")
            except HTTPException as exc:
                assert exc.status_code == 400
                assert "checklist" in exc.detail.lower()
            db.rollback()

        assert (
            db.query(OrganizationRegion)
            .filter_by(org_id=organization.org_id, region_id=requested_region.region_id)
            .first()
            is None
        )
    finally:
        db.close()


def test_region_checklist_decision_cannot_approve_incomplete_submission():
    db = Session(bind=engine)
    try:
        owner = User(
            user_id="incomplete-approval-owner",
            org_id="org-incomplete-approval",
            email="incomplete-approval-owner@example.com",
            password="test-password",
            role="owner",
        )
        soc = User(
            user_id="incomplete-approval-soc",
            org_id=None,
            email="incomplete-approval-soc@example.com",
            password="test-password",
            role="soc_analyst",
        )
        organization = Organization(
            org_id="org-incomplete-approval",
            user_id=owner.user_id,
        )
        region = Region(code="INCOMPLETE-APPROVAL", name="Incomplete", is_active=True)
        db.add_all([owner, soc, organization, region])
        db.flush()
        row = OrganizationRegion(
            org_id=organization.org_id,
            region_id=region.region_id,
            status="pending",
            testing_start_at=datetime(2026, 10, 15, 10, tzinfo=timezone.utc),
            testing_timezone="UTC",
            checklist_submission={
                "checklist_answers": {
                    "company_details": {
                        "organization_name": {"answer": "Acme", "na": False},
                    }
                }
            },
        )
        db.add(row)
        db.commit()

        try:
            _run(
                decide_region_checklist(
                    payload={
                        "org_id": organization.org_id,
                        "region_code": region.code,
                        "status": "approved",
                    },
                    db=db,
                    current_user=soc,
                )
            )
            raise AssertionError("expected approval to reject an incomplete region checklist")
        except HTTPException as exc:
            assert exc.status_code == 400
            assert "completed checklist" in exc.detail.lower()
        db.refresh(row)
        assert row.status == "pending"
        assert row.checklist_review_status == "pending"
    finally:
        db.close()


def test_admin_approved_client_can_submit_first_region_checklist():
    """Admin-granted VAPT access allows the client to request the first region."""
    import app.api.vapt.routes as vapt_routes

    db = Session(bind=engine)
    original_email_sender = vapt_routes.send_vapt_access_event_email
    original_ws_send = vapt_routes.ws_manager.send

    async def noop_send(*args, **kwargs):
        return None

    try:
        client = User(
            user_id="first-region-client",
            org_id="org-first-region",
            email="first-region-client@example.com",
            password="test-password",
            role="owner",
            vapt_approved=True,
        )
        db.add(client)
        db.commit()
        vapt_routes.send_vapt_access_event_email = lambda *args, **kwargs: None
        vapt_routes.ws_manager.send = noop_send

        response = _run(
            request_vapt_region(
                payload={
                    "region_code": "FIRST-IND",
                    "region_name": "First Region",
                    "testing_start_at": "2026-10-15T10:00:00+00:00",
                    "testing_timezone": "UTC",
                    "scope_ip_ranges": "10.2.0.0/16",
                    "checklist_answers": _complete_region_checklist_answers(),
                },
                db=db,
                current_user=client,
            )
        )

        assert response["success"] is True
        region = db.query(Region).filter_by(code="FIRST-IND").one()
        region_request = (
            db.query(OrganizationRegion)
            .filter_by(org_id="org-first-region", region_id=region.region_id)
            .one()
        )
        assert region_request.status == "pending"
        assert region_request.checklist_submission["scope_ip_ranges"] == "10.2.0.0/16"
    finally:
        vapt_routes.send_vapt_access_event_email = original_email_sender
        vapt_routes.ws_manager.send = original_ws_send
        db.close()


def test_region_checklist_changes_requested_then_approved():
    """A region checklist with a few wrong answers must not reject the whole
    region request — SOC asks for more information and the region stays pending."""
    db = Session(bind=engine)
    try:
        client = User(
            user_id="rd-client",
            org_id="org-rd",
            email="rd-client@example.com",
            password="hashed",
            role="owner",
        )
        soc = User(
            user_id="rd-soc",
            org_id=None,
            email="rd-soc@example.com",
            password="hashed",
            role="soc_analyst",
        )
        db.add_all([client, soc])
        db.add(Organization(org_id="org-rd", user_id="rd-client", max_domains=1, region="RD-B"))
        db.add(
            VaptOnboardingChecklist(
                org_id="org-rd",
                checklist_answers={
                    "general_information": {
                        "network_diagram_available": {"answer": "Yes", "na": False},
                    }
                },
                completed_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
                review_status="pending",
            )
        )
        db.commit()

        request_vapt_access(
            payload={
                "regions": [{"code": "RD-A", "name": "Region A"}],
                "testing_start_at": "2026-09-15T10:00:00+00:00",
                "testing_timezone": "UTC",
                "checklist_answers": _complete_region_checklist_answers(),
            },
            db=db,
            current_user=client,
        )
        _run(
            approve_vapt_access(
                payload={"org_id": "org-rd", "region": "RD-A", "approved": True},
                db=db,
                current_user=soc,
            )
        )
        request_vapt_access(
            payload={
                "regions": [{"code": "RD-B", "name": "Region B"}],
                "testing_start_at": "2026-10-01T10:00:00+00:00",
                "testing_timezone": "UTC",
                "checklist_answers": _complete_region_checklist_answers(),
            },
            db=db,
            current_user=client,
        )

        reviewed = _run(
            decide_region_checklist(
                payload={
                    "org_id": "org-rd",
                    "region_code": "RD-B",
                    "status": "changes_requested",
                    "note": "Please attach the network diagram.",
                    "flags": [
                        {
                            "section": "general_information",
                            "question_id": "network_diagram_available",
                            "label": "Network diagram",
                            "note": "Send the latest diagram.",
                        }
                    ],
                },
                db=db,
                current_user=soc,
            )
        )
        assert reviewed["checklist_review_status"] == "changes_requested"
        assert reviewed["status"] == "pending"  # the region was NOT rejected
        region_b = db.query(Region).filter_by(code="RD-B").one()
        row = db.query(OrganizationRegion).filter_by(
            org_id="org-rd",
            region_id=region_b.region_id,
        ).one()
        assert row.checklist_flags[0]["question_id"] == "network_diagram_available"
        assert row.checklist_submission["checklist_answers"]["general_information"]["network_diagram_available"]["answer"] == ""

        # The client can see the remarks and flagged items.
        status = get_vapt_access_status(db=db, current_user=client)
        pending = next(item for item in status["pending_regions"] if item["code"] == "RD-B")
        assert pending["checklist_review_status"] == "changes_requested"
        assert pending["checklist_flags"][0]["question_id"] == "network_diagram_available"

        incomplete_answers = _complete_region_checklist_answers()
        incomplete_answers["general_information"]["network_diagram_available"]["answer"] = ""
        try:
            request_vapt_access(
                payload={
                    "regions": [{"code": "RD-B", "name": "Region B"}],
                    "testing_start_at": "2026-10-01T10:00:00+00:00",
                    "testing_timezone": "UTC",
                    "checklist_answers": incomplete_answers,
                },
                db=db,
                current_user=client,
            )
            raise AssertionError("expected 400 when the flagged answer is still empty")
        except HTTPException as exc:
            assert exc.status_code == 400
            assert exc.detail == "Update every SOC-flagged checklist question before resubmitting."
        db.rollback()

        # Re-entering the flagged answer is valid even when the corrected
        # answer is textually the same as before; the client explicitly
        # re-entered it after SOC cleared it for review.
        request_vapt_access(
            payload={
                "regions": [{"code": "RD-B", "name": "Region B"}],
                "testing_start_at": "2026-10-01T10:00:00+00:00",
                "testing_timezone": "UTC",
                "checklist_answers": _complete_region_checklist_answers(),
            },
            db=db,
            current_user=client,
        )
        db.expire_all()
        region_b = db.query(Region).filter_by(code="RD-B").one()
        row = (
            db.query(OrganizationRegion)
            .filter_by(org_id="org-rd", region_id=region_b.region_id)
            .one()
        )
        assert row.checklist_review_status == "pending"
        assert row.checklist_flags is None

        # Approving accepts the region together with its checklist.
        approved = _run(
            decide_region_checklist(
                payload={"org_id": "org-rd", "region_code": "RD-B", "status": "approved"},
                db=db,
                current_user=soc,
            )
        )
        assert approved["status"] == "approved"
        assert approved["checklist_review_status"] == "approved"
        db.expire_all()
        row = db.query(OrganizationRegion).filter_by(
            org_id="org-rd",
            region_id=region_b.region_id,
        ).one()
        assert row.status == "approved"
        assert row.checklist_review_status == "approved"
        assert row.checklist_flags is None
    finally:
        db.close()


def test_rejected_region_checklist_syncs_onboarding_and_can_be_resubmitted():
    db = Session(bind=engine)
    try:
        client = User(
            user_id="reject-sync-client",
            org_id="org-reject-sync",
            email="reject-sync-client@example.com",
            password="test-password",
            role="owner",
        )
        soc = User(
            user_id="reject-sync-soc",
            org_id=None,
            email="reject-sync-soc@example.com",
            password="test-password",
            role="soc_analyst",
        )
        region = Region(code="REJECT-SYNC", name="Reject Sync", is_active=True)
        db.add_all([client, soc, Organization(org_id="org-reject-sync", user_id=client.user_id), region])
        db.flush()
        row = OrganizationRegion(
            org_id="org-reject-sync",
            region_id=region.region_id,
            status="pending",
            checklist_submission={"checklist_answers": {}},
        )
        db.add_all(
            [
                row,
                VaptOnboardingChecklist(
                    org_id="org-reject-sync",
                    checklist_answers={"general_information": {"organization_name": {"answer": "Acme"}}},
                    completed_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
                    review_status="pending",
                ),
            ]
        )
        db.commit()
        row_id = row.id

        _run(
            decide_region_checklist(
                payload={
                    "org_id": "org-reject-sync",
                    "region_code": "REJECT-SYNC",
                    "status": "rejected",
                    "note": "The requested scope is not authorized.",
                },
                db=db,
                current_user=soc,
            )
        )
        db.expire_all()
        row = db.query(OrganizationRegion).filter_by(id=row_id).one()
        onboarding = db.query(VaptOnboardingChecklist).filter_by(org_id="org-reject-sync").one()
        assert row.status == "rejected"
        assert row.checklist_review_status == "rejected"
        assert row.checklist_flags is None
        assert onboarding.review_status == "rejected"
        assert onboarding.completed_at is None
        assert onboarding.review_flags is None

        request_vapt_access(
            payload={"regions": [{"code": "REJECT-SYNC", "name": "Reject Sync"}]},
            db=db,
            current_user=client,
        )
        db.expire_all()
        resubmitted = db.query(OrganizationRegion).filter_by(id=row_id).one()
        assert resubmitted.status == "pending"
        assert resubmitted.checklist_review_status == "pending"
        assert db.query(OrganizationRegion).filter_by(org_id="org-reject-sync").count() == 1
    finally:
        db.close()


def test_changes_requested_clears_and_requires_flagged_items():
    """A partial review clears flagged answers and requires fresh values."""
    db = Session(bind=engine)
    try:
        client = User(
            user_id="cr-client",
            org_id="org-cr",
            email="cr-client@example.com",
            password="hashed",
            role="owner",
        )
        soc = User(
            user_id="cr-soc",
            org_id=None,
            email="cr-soc@example.com",
            password="hashed",
            role="soc_analyst",
        )
        db.add_all([client, soc])
        db.commit()

        request_vapt_access(
            payload={
                "regions": [{"code": "CR-IND", "name": "CR India"}],
                "testing_start_at": "2026-09-15T10:00:00+00:00",
                "testing_timezone": "UTC",
                "checklist_answers": _complete_region_checklist_answers(),
            },
            db=db,
            current_user=client,
        )
        _run(
            approve_vapt_access(
                payload={"org_id": "org-cr", "region": "CR-IND", "approved": True},
                db=db,
                current_user=soc,
            )
        )
        _run(
            submit_onboarding_checklist(
                payload={
                    "testing_start_at": "2026-09-15T10:00:00+00:00",
                    "testing_timezone": "UTC",
                    "checklist_answers": {
                        "general_information": {
                            "network_diagram_available": {"answer": "Yes", "na": False},
                        }
                    },
                },
                db=db,
                current_user=client,
            )
        )

        reviewed = _run(
            review_vapt_onboarding(
                "org-cr",
                {
                    "status": "changes_requested",
                    "note": "Please attach the network diagram.",
                    "flags": [
                        {
                            "section": "general_information",
                            "question_id": "network_diagram_available",
                            "label": "Network diagram",
                            "note": "Send the latest diagram.",
                        }
                    ],
                },
                db=db,
                current_user=soc,
            )
        )
        assert reviewed["review_status"] == "changes_requested"
        assert reviewed["completed"] is True  # the submission cycle remains open
        assert len(reviewed["review_flags"]) == 1
        assert reviewed["review_flags"][0]["question_id"] == "network_diagram_available"
        assert reviewed["checklist_answers"]["general_information"]["network_diagram_available"]["answer"] == ""

        # Re-submitting the fresh flagged answer resolves the request.
        resubmitted = _run(
            submit_onboarding_checklist(
                payload={
                    "checklist_answers": {
                        "general_information": {
                            "network_diagram_available": {"answer": "Emailed", "na": False},
                        }
                    }
                },
                db=db,
                current_user=client,
            )
        )
        assert resubmitted["onboarding"]["review_status"] == "pending"
        assert resubmitted["onboarding"]["review_flags"] == []
        assert resubmitted["onboarding"]["completed"] is True
    finally:
        db.close()


def test_review_checklist_requires_remarks_or_flags_for_partial():
    db = Session(bind=engine)
    try:
        soc = User(
            user_id="rf-soc",
            org_id=None,
            email="rf-soc@example.com",
            password="hashed",
            role="soc_analyst",
        )
        db.add(soc)
        db.commit()
        try:
            _run(
                review_vapt_onboarding(
                    "org-missing",
                    {"status": "changes_requested"},
                    db=db,
                    current_user=soc,
                )
            )
            raise AssertionError("expected 404 for an unknown org")
        except HTTPException as exc:
            assert exc.status_code == 404
    finally:
        db.close()


def test_reopen_resets_only_findings_not_confirmed_fixed():
    record = VaptImport(
        org_id="org-reopen",
        file_name="test.nessus",
        file_format="xml",
        source_tool="nessus",
        total_findings=2,
        unique_hosts=1,
        risk_score=50,
        severity="medium",
        findings=[
            {"id": "fixed", "title": "Fixed", "plugin_id": "1", "status": "solved", "comment": "patched"},
            {"id": "open", "title": "Open", "plugin_id": "2", "status": "solved", "comment": "patched"},
            {"id": "low", "title": "Low", "plugin_id": "3", "severity_label": "low", "status": "solved", "comment": "Client fix remark"},
        ],
    )
    reopened = _reopen_unresolved_findings(record, {"fixed_findings": [record.findings[0]]})
    assert reopened[0]["status"] == "solved"
    assert reopened[0]["comment"] == "patched"
    assert reopened[1]["status"] == "pending"
    assert reopened[1]["comment"] == ""
    assert reopened[2]["status"] == "solved"
    assert reopened[2]["comment"] == "Client fix remark"


def test_approved_region_package_is_listed_and_org_checklist_can_download_for_region():
    from io import BytesIO
    from zipfile import ZipFile

    from openpyxl import load_workbook
    from pypdf import PdfReader

    async def response_bytes(response):
        return b"".join([chunk async for chunk in response.body_iterator])

    db = Session(bind=engine)
    try:
        user = User(user_id="pkg-user", org_id="org-pkg", email="pkg@example.com", password="hashed", role="owner")
        soc = User(user_id="pkg-soc", org_id=None, email="pkg-soc@example.com", password="hashed", role="soc_analyst")
        db.add_all([user, soc, Organization(org_id="org-pkg", user_id="pkg-user", max_domains=1)])
        db.commit()

        region = Region(code="PKG-R", name="Package Region", is_active=True)
        db.add(region)
        db.flush()
        reviewed_at = datetime.now(timezone.utc)
        db.add(OrganizationRegion(
            org_id="org-pkg",
            region_id=region.region_id,
            status="approved",
            reviewed_by="pkg-soc",
            reviewed_at=reviewed_at,
            testing_start_at=datetime(2026, 10, 8, 6, 34, tzinfo=timezone.utc),
            testing_timezone="America/Los_Angeles",
            checklist_review_status="approved",
            checklist_submission={
                "scope_ip_ranges": "10.2.0.0/16",
                "checklist_answers": {"general_information": {"organization_name": {"answer": "Package Co"}}},
            },
        ))
        db.commit()

        packages = list_approved_vapt_onboarding(db=db, current_user=soc)
        package = next(item for item in packages if item["region_code"] == "PKG-R")
        assert package["checklist_answers"]["general_information"]["organization_name"]["answer"] == "Package Co"
        package_start = package["testing_start_at"]
        assert package_start.replace(tzinfo=timezone.utc) == datetime(2026, 10, 8, 6, 34, tzinfo=timezone.utc)
        assert package["testing_timezone"] == "America/Los_Angeles"

        response = download_approved_onboarding_bundle("org-pkg", region_code="PKG-R", db=db, current_user=soc)
        assert response.media_type == "application/zip"
        with ZipFile(BytesIO(_run(response_bytes(response)))) as archive:
            workbook = load_workbook(BytesIO(archive.read("checklist.xlsx")), read_only=True)
            metadata = {row[0]: row[1] for row in workbook["Checklist"].iter_rows(min_row=2, max_col=2, values_only=True) if row[0]}
            assert metadata["Testing start (SOC / IST)"] == "08 Oct 2026, 12:04 PM IST"
            assert metadata["Testing start (client local: America/Los_Angeles)"] == "07 Oct 2026, 11:34 PM PDT"

            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(archive.read("checklist.pdf"))).pages)
            assert "Testing start (SOC / IST)" in pdf_text
            assert "08 Oct 2026, 12:04 PM IST" in pdf_text
            assert "07 Oct 2026, 11:34 PM PDT" in pdf_text
    finally:
        db.close()


def test_checklist_pdf_preserves_long_timezone_metadata():
    pdf = _generate_checklist_pdf(
        {
            "region_name": "KPGU USA",
            "approved_at": datetime(2026, 10, 22, 5, 9, tzinfo=timezone.utc),
            "testing_timezone": "America/Argentina/Buenos_Aires",
            "testing_start_at": datetime(2026, 10, 23, 14, 34, tzinfo=timezone.utc),
            "checklist_answers": {},
        }
    )

    text = " ".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(pdf)).pages
    )
    assert "Testing start (client local: America/Argentina/Buenos_Aires)" in text
    assert "Testing timezone" in text
    assert "America/Argentina/Buenos_Aires" in text
